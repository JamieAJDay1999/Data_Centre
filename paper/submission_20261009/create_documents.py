"""Create the editable submission correspondence from checked manuscript content."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
PAPER = OUT.parent
TITLE = 'Characterisation and Quantification of Data Centre Flexibility for Power System Support'
AUTHORS = 'James Day, Mehmet Türker Takci and Meysam Qadrdan'
DATE = '13 October 2026'

def document(title, footer=True):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size in [('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        st=doc.styles[name]
        st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)
    doc.core_properties.author=AUTHORS
    doc.core_properties.title=title
    if footer:
        p=sec.footer.paragraphs[0]
        p.alignment=2
        r=p.add_run('APEN-D-25-21553 resubmission  |  '); r.font.size=Pt(9)
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)
    return doc

def paragraphs(doc, text):
    for value in text.strip().split('\n\n'):
        doc.add_paragraph(value.strip())

doc=document('Cover letter',False)
paragraphs(doc, DATE + '''

Professor Zita Vale\nEditor-in-Chief\nApplied Energy

Dear Professor Vale,

Please consider our substantially revised manuscript, “'''+TITLE+'''”, for publication as a research article in Applied Energy. This is a resubmission following rejection of manuscript APEN-D-25-21553. We refer explicitly to that earlier submission in accordance with your decision letter. The author order differs from the previous submission: James Day is now first author, followed by Mehmet Türker Takci and Meysam Qadrdan. All authors have agreed to this order.

The study develops a whole-facility framework for quantifying the magnitude and feasible duration of data centre flexibility through coordinated IT scheduling, UPS batteries and cooling. Its contribution is the assessment of grid-facing responses from a common operating baseline, with operational constraints and comparative post-event recovery conditions. This connects energy management within data centres to the practical characterisation of demand-side flexibility for power system support.

The revision substantially expands the evidence from a single-day cost comparison to linked daily optimisation over 2025, with fourteen annual scenario chains and sensitivities to workload flexibility and storage capacity. The central annual electricity cost saving is 5.38%; this supersedes the 10.02% single-day estimate in the previous submission and is not a like-for-like comparison with it. We provide a component cost breakdown, an explicit optimisation workflow, clearer workload-input provenance, detailed solver and numerical checks, and revised figures. We also distinguish certified feasible durations from unproven maximum durations and clarify the limits of the network and economic interpretations.

We enclose a clean revised manuscript, a copy with changes marked against the rejected version, and point-by-point responses to all comments from Reviewers 2 and 5. The submission also includes highlights, supplementary information and the declaration of interests.

All authors have approved this version and its submission. The work is original, has not been published previously and is not under consideration elsewhere. The authors declare no competing interests. Funding, author contributions and the use of ChatGPT Codex and Claude Code for manuscript preparation and coding are disclosed in the manuscript. I remain the corresponding author.

Thank you for considering this revised submission.

Yours sincerely,

James Day\nOn behalf of all authors\nCardiff University\nDayJa1@cardiff.ac.uk''')
doc.save(OUT/'cover_letter.docx')

doc=document('Highlights',False)
doc.add_heading('Highlights',0)
doc.add_paragraph(TITLE)
highlights=(OUT/'highlights.txt').read_text(encoding='utf8').splitlines()
for text in highlights:
    text=text.strip().lstrip('- ').strip()
    if text:
        assert len(text)<=85,(len(text),text)
        doc.add_paragraph(text,style='List Bullet')
doc.save(OUT/'highlights.docx')

# Preserve reviewers' supplied wording, excluding the colleague's informal notes.
raw=(PAPER/'revisions/reviewers_comments_and_colleague_response.md').read_text(encoding='utf8')
r2,r5=raw.split('## Reviewer #2',1)[1].split('## Reviewer #5',1)
def quoted_comments(part, implicit_first=False):
    values=[]
    for line in part.splitlines():
        if not line.startswith('> '): continue
        text=line[2:].strip()
        number=re.match(r'^(\d+)\.\s*(.*)',text)
        if number:
            values.append(number[2])
        elif not values:
            values.append(text)
        else:
            values[-1]+=' '+text
    return values
comments={2:quoted_comments(r2),5:quoted_comments(r5,True)}
assert len(comments[2])==7 and len(comments[5])==8

responses={
2:[
('Novelty and the role of the UPS model',
'''We agree that the UPS formulation is a standard storage model and should not be presented as a new battery model. We have clarified that the contribution lies in the integrated, duration-aware flexibility assessment. IT arrivals and delay windows, UPS energy, thermal states and TES energy jointly constrain the deviation that the facility can sustain from its cost-minimising baseline. Scenario 3 evaluates start time, requested power deviation and service duration while enforcing comparative recovery conditions.

The revised introduction and literature review distinguish this contribution from asset-level scheduling and earlier coordinated energy-management studies. The discussion explains the cross-asset interactions revealed by the model, while expressly avoiding a claim that the present results measure an integration premium over separately optimised subsystems; that would require an additional comparator. The UPS model remains conventional so that the contribution and its limits are explicit.''',
'Introduction, pp. 1 and 4–5; Sections 2–3, pp. 5–10; Section 4.3, p. 13; Section 5.5, pp. 18–19.'),
('Figure clarity, terminology and presentation',
'''We have revised the figure presentation and the accompanying interpretation. The upward-flexibility and downward-flexibility discussions are now separate and identify the sign convention explicitly: negative ΔP denotes a load reduction, while positive ΔP denotes a load increase. The component plots carry panel labels (a)–(d), and their captions identify the start times and requested magnitudes. All plots are now drawn at their printed size, so figure text is 7–8 pt and consistent with the main text; Figures 2 and 3 have been redrawn in the same style. The plots use consistent colours and are supplied in vector form, and the multi-panel component plots span the page width.

The power/price and IT-workload panels retain their complementary roles but now share a common time axis and consistent presentation. To keep terminology uniform, the abbreviation “DC” is no longer used: “data centre” is written in full throughout the text, captions and figures. Inserting the workflow figure changes the numbering: original Fig. 4 corresponds to revised Fig. 5, and the upward/downward component plots are now Figs. 8 and 9.''',
'Figs. 2–3 and 5–10, pp. 7–8 and 14–18; Section 5.2, pp. 14–17; terminology throughout.'),
('Case-study structure and additional sensitivities',
'''We have reorganised the case-study explanation around the three-scenario workflow and substantially expanded the numerical evidence. The annual comparison now uses linked daily optimisation over all of 2025. Fourteen completed annual chains cover the reference case, the central coordinated case and one-at-a-time changes in flexible workload share, UPS energy capacity and TES energy capacity at multipliers 0.5, 0.75, 1.25 and 1.5. The flexible-share construction preserves total arriving work and clips the flexible portion at the total workload.

The full year includes the observed variation in signed reference prices, including negative prices. We have not added a controlled price-volatility sweep: annual exposure to realised prices does not isolate the causal effect of volatility. The revised text states this limit explicitly, so the sensitivity results support conclusions about the tested workload and equipment assumptions, not universality across price regimes or sites. The full endpoint table and solver exceptions are supplied in the supplement.''',
'Section 4 and Fig. 4, pp. 10–13; Sections 4.4 and 5.4, pp. 13 and 17–18; Fig. 10, p. 18; Supplementary Sections S4–S5.'),
('Cost composition and the reported saving',
'''We have added a full annual signed electricity-cost breakdown. The central annual reference cost is GBP 550,503.59 and the coordinated cost is GBP 520,876.68, giving GBP 29,626.91 (5.38%) savings. IT electricity expenditure falls by GBP 16,722.17 and direct CRAC expenditure by GBP 14,425.46; TES charging adds GBP 6,897.29, net UPS exchange reduces expenditure by GBP 5,376.57, and auxiliary expenditure is unchanged. Table VI reconciles these terms.

The revised annual result replaces the 10.02% saving previously reported for a single day. The annual 5.38% covers a different assessment period and should not be interpreted as a direct percentage-point comparison with the single-day result. These components are accounting terms within a coordinated dispatch, not independent asset profits. The text also reports the increase in annual electricity use and peak import and excludes service revenues and degradation costs from the valuation.''',
'Section 5.3, p. 17; Table VI, p. 18; Conclusion, p. 19; Supplementary Section S4.'),
('Problem class, implementation and reproducibility',
'''The revised manuscript identifies the solved model as a mixed-integer linear programme implemented in Pyomo with HiGHS. Binary variables enforce storage modes and the IT power curve uses a four-segment piecewise-linear DLOG representation with non-uniform breakpoints. The manuscript specifies this approximation and its maximum interpolation error (6.241 kW, or 0.624% of the 1 MW IT rating), the optimisation objectives, temporal boundaries, state/workload carry-over and event recovery criteria.

The supplement records the retained software versions, hardware, time limits, requested and accepted gaps, and physical/accounting checks. It distinguishes feasible time-limited incumbents from proven optimal solutions. Three annual sensitivity horizons exceed the normal 1% acceptance threshold and are disclosed with their small absolute daily bound widths; these are not hidden by a blanket optimality claim. All 148 reported positive-duration event dispatches meet the separate event gap criterion. The adopted inputs and parameters are reported in full in the manuscript and supplement, and the electricity prices are publicly available.''',
'Sections 4.2–4.3, pp. 12–13; Fig. 4, p. 11; Supplementary Sections S2–S5; Data availability, p. 19.'),
('Grid connection and network implications',
'''We agree that cost-minimising operation can increase the network peak and that this affects the interpretation of flexibility. The revised results explicitly report an increase in annual peak import from 974.9 to 1,682.5 kW and annual grid electricity from 6.718 to 6.789 GWh, despite the lower reference-price expenditure.

The model assesses feasible facility responses at the connection point; it does not contain an upstream power-flow or dynamic-stability model. We therefore do not claim that the optimised load profile is automatically admissible on a particular network or improves stability. The discussion identifies topology, voltage and thermal limits and other concurrent loads as necessary inputs for a subsequent network assessment. The reported responses are candidate flexibility offers under the stated facility assumptions, to be screened against those network constraints and service requirements.''',
'Section 5.3, p. 17; Section 5.5, pp. 18–19; Conclusion, p. 19.'),
('Workload data sources and representativeness',
'''We have made the status and construction of the workload inputs explicit. They are stylised adopted profiles informed by the cited workload literature, rather than measured site data, a statistically representative facility sample or a deterministic reconstruction of a published trace. Table II summarises the evidence informing the assumptions, while Tables III–IV and Supplementary Table S1 expose the numerical inputs and their interpretation.

The model uses the 96 quarter-hour utilisation values directly. The main-paper workload plot/table displays hourly arithmetic means, not the first quarter-hour in each hour. The 24 hourly tranche-share rows are expanded by repeating each row for its four consecutive quarter-hours. The exact flexible/inflexible arrivals and tranche shares are tabulated in Supplementary Table S1, and the supplement explains the delay-window convention and clipping rule used in the sensitivity analysis. We have not represented a network electricity-demand dataset as measured CPU utilisation. The revised limitations and sensitivity results qualify the transferability of these adopted profiles.''',
'Section 3.1, pp. 6–8; Tables II–IV and Fig. 2; Section 5.4 and Section 5.5 limitations; Supplementary Section S1 and Table S1.')
],
5:[
('Aggregate modelling rationale and distinct contribution',
'''We have clarified both the modelling purpose and the variables retained. The facility-level model preserves arriving computational work and eligible execution windows, UPS and TES energy, storage power limits and mode exclusivity, and the thermal states that couple IT activity to cooling demand. These are the quantities that constrain a grid-facing power deviation and its duration. Aggregation avoids modelling individual instructions, rack airflow and task dependencies, making repeated horizon and duration evaluations tractable.

This abstraction does not establish task-level quality of service for arbitrary applications. The text now explains that CPU-hours are a processor-occupancy representation and states the limitations of the adopted workloads, fixed environmental assumptions and quarter-hour resolution. The duration assessment uses a common operating baseline and explicit recovery conditions across the coupled resources. The integrated-dispatch discussion illustrates those mechanisms without claiming a measured integration premium against a separately optimised comparator.''',
'Introduction and Sections 2–3, pp. 1 and 4–10; Sections 4.2–4.3, pp. 12–13; Section 5.5, pp. 18–19; Supplementary Sections S1–S3.'),
('Introduction structure and research focus',
'''The revised introduction opens with the expansion and concentrated electricity demand of data centres, connects this directly to power-system flexibility, and sets out the facility resources relevant to the study. The research gap and contributions distinguish coordinated energy scheduling from quantifying feasible power-deviation durations relative to a common operating baseline. The revised abstract likewise states the integrated framework, annual baseline, duration assessment and principal findings directly.''',
'Abstract and Introduction, pp. 1 and 4–5; Section 2, pp. 5–6.'),
('Transparency of workload profiles and ratios',
'''We have expanded the workload provenance and construction description, as also detailed in our response to Reviewer 2, comment 7. The profiles are explicitly described as stylised assumptions informed by literature, with no claim that the chosen flexible/inflexible ratio is universal or directly measured. Table II identifies the supporting literature and its role; Tables III–IV and Supplementary Table S1 give the adopted quantities.

The supplement specifies the 96 quarter-hour arrays, the arithmetic averaging used for hourly presentation, and the repetition of each hourly tranche row across four consecutive intervals. It explains arrival/tranche work, final eligible execution intervals and the sensitivity clipping rule that preserves total work. The complete quarter-hour inputs are tabulated. These changes make the calculation reproducible while separating the supporting literature from the authors’ numerical assumptions.''',
'Section 3.1, pp. 6–8; Tables II–IV; Supplementary Section S1 and Table S1.'),
('Justification of utilisation multiplied by duration',
'''We have added a discussion with references to Dayarathna et al., the Google cluster-trace analysis of Rasheduzzaman et al. and Cao et al. The product of CPU utilisation and execution duration is interpreted as accumulated processor occupancy at the aggregate facility level. For example, 50% utilisation for two hours and 100% for one hour represent equal CPU-hour demand in this abstraction.

This equivalence is not a claim that heterogeneous real jobs have identical instruction throughput or can always trade utilisation for runtime without other constraints. The manuscript limits its scope to aggregate CPU-work scheduling, retains eligible execution windows and conservation of work, and identifies task dependencies and GPU-specific controls as outside the representation.''',
'Section 3.1.1 and Eq. (2), p. 8; Section 3.1.2, pp. 8–9; Section 5.5 limitations, p. 19.'),
('Optimisation workflow',
'''We have added a three-scenario workflow diagram and clarified the inputs, objectives and information flow. Scenario 1 establishes a workload-at-arrival reference with storage disabled; cooling is still determined consistently with the facility constraints. Scenario 2 co-optimises operation against signed reference prices over linked daily horizons, committing each day and passing physical states and outstanding work to the next.

Scenario 3 starts from the resulting undisturbed trajectory. For each requested start time and power deviation, the duration search solves event-plus-recovery dispatch problems and applies the recorded feasibility and solution-quality checks. The output is a grid of certified feasible durations and corresponding component dispatches. The revised explanation distinguishes timeouts from infeasibility and notes that assessment caps and a moving recovery reference prevent interpreting every reported duration as a proven global maximum.''',
'Fig. 4, p. 11; Sections 4.1–4.3, pp. 11–13; Supplementary Sections S2–S3.'),
('Selection and meaning of component decompositions',
'''We have clarified that each decomposition is an accepted feasible MILP dispatch for the specified start time, power deviation and certified event duration. The event-plus-recovery signed electricity-cost objective selects among feasible dispatches, subject to the recorded solver gap. It is therefore one economically preferred realisation; the model does not establish a unique physical attribution or average over all feasible optima.

The plotted components are differences from the same undisturbed trajectory and sum to the metered grid-power deviation. Opposing component responses are retained because some assets may absorb energy while others over-deliver the requested net change. Recovery is assessed through terminal storage, temperature and outstanding-work conditions. The figures distinguish the event window from the unplotted recovery period and mark unresolved adjacent-duration trials. The independent component accounting check covers all 148 positive-duration dispatches.''',
'Section 4.3, p. 13; Section 5.2 and Figs. 8–9, pp. 14–17; Supplementary Sections S3 and S5.'),
('Insights from modelling the subsystems together',
'''We have added a dedicated discussion of the value of integrated co-optimisation. It explains the coupled reduction in IT heat and cooling demand when work is deferred, the change from UPS-dominated morning reductions to workload/cooling-dominated afternoon reductions, and the interaction of TES charging and direct cooling with the other assets. The component plots show simultaneous opposing responses that meet a single connection-point target.

These mechanisms demonstrate why installed capacity alone does not determine feasible duration: the baseline workload, storage energy and thermal state matter together. For example, many 100–200 kW reductions end at 07:00 or 18:00 because the undisturbed schedule already supplies the whole IT load from the UPS and has deferred all eligible work at those times. We also state the evidential limit explicitly. The present study reveals coupled dispatch mechanisms but does not quantify an integration premium over independently optimised subsystems, because that comparator was not run.''',
'Section 5.2 and Figs. 8–9, pp. 14–17; Section 5.5, pp. 18–19.'),
('Distinguishing charging and discharging in the resource plot',
'''We have revised the resource-power plot to distinguish positive consumption/charging from negative UPS discharge. UPS discharge now uses a hatch pattern and a separate legend entry from charging. The caption explains the sign convention and clarifies that TES discharge reduces chiller electricity demand rather than injecting electricity. The plot is supplied as a full-width vector figure for readability. The insertion of the workflow diagram means the original Fig. 5 is now Fig. 6.''',
'Revised Fig. 6 and its caption, p. 15.')
]
}

doc=document('Response to reviewers')
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Heading 3'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 3'].paragraph_format.space_after=Pt(3)
doc.add_heading('Response to the editor and reviewers',0)
doc.add_paragraph(TITLE)
doc.add_paragraph('Previous manuscript: APEN-D-25-21553\n'+AUTHORS+'\n'+DATE)
paragraphs(doc, '''Dear Editor and Reviewers,

Thank you for the detailed comments. We have substantially revised the manuscript and provide a point-by-point response below. The revision includes full-year operation and resource sensitivities, an explicit optimisation workflow, transparent workload inputs, component cost accounting, numerical checks, clearer figures and more precise limits on the flexibility and network claims.

In accordance with the decision letter, the cover letter identifies the previous manuscript, and we provide clean and marked revised copies. The marked copy compares the revised manuscript with the version submitted as APEN-D-25-21553. Blue text identifies additions or replacements; deleted material is suppressed. Page numbers below refer to the clean revised manuscript. Figure numbers have changed after insertion of the workflow diagram. Reviewer comments are reproduced in full.''')
for reviewer in [2,5]:
    doc.add_heading(f'Reviewer {reviewer}',1)
    for n,(quote,(title,answer,location)) in enumerate(zip(comments[reviewer],responses[reviewer]),1):
        doc.add_heading(f'Comment {reviewer}.{n} — {title}',2)
        p=doc.add_paragraph()
        p.paragraph_format.keep_with_next=True
        r=p.add_run('Reviewer: '); r.bold=True
        r=p.add_run(quote); r.italic=True; r.font.size=Pt(10)
        doc.add_paragraph('Response',style='Heading 3')
        paragraphs(doc,answer)
        p=doc.add_paragraph()
        p.add_run('Location: ').bold=True
        p.add_run(location)
doc.save(OUT/'response_to_reviewers.docx')

doc=document('Submission metadata and declarations')
doc.add_heading('Submission metadata and declarations',0)
doc.add_paragraph('Copy-ready information for the Applied Energy submission form. Previous manuscript: APEN-D-25-21553.')
for heading,text in [
('Article title',TITLE),
('Article type','Full Length Research Article.'),
('Authors and affiliation','1. James Day — corresponding author\n2. Mehmet Türker Takci\n3. Meysam Qadrdan\n\nAll authors: Cardiff University, Queen’s Buildings, The Parade, Cardiff CF24 3AA, United Kingdom.\nCorrespondence: DayJa1@cardiff.ac.uk'),
('Keywords','Data centre; power system flexibility; demand side flexibility; load shifting; thermal inertia'),
('Previous submission','This is a substantially revised resubmission of APEN-D-25-21553, previously rejected by Applied Energy. The previous decision letter permits resubmission after significant revision. Responses address the reports of Reviewers 2 and 5. The previous submission listed Mehmet Türker Takci first; James Day is now first author.'),
('Funding','Engineering and Physical Sciences Research Council (EPSRC) and Economic and Social Research Council (ESRC), through the Energy Demand Research Centre, grant EP/Y010078/1. The funders had no role in study design, analysis, manuscript preparation or the decision to submit.'),
('Competing interests','The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.'),
('CRediT authorship contribution statement','James Day: Writing – original draft, Investigation, Methodology, Conceptualization.\nMehmet Türker Takci: Writing – original draft, Investigation, Methodology, Conceptualization.\nMeysam Qadrdan: Writing – review & editing, Supervision, Conceptualization.'),
('Generative AI declaration','During the preparation of this work, the authors used ChatGPT Codex (OpenAI) and Claude Code (Anthropic) to assist with manuscript preparation, editing and coding, including the Python code used to plot figures from the supplied numerical inputs and results, the workflow diagram and the redrawn facility schematic, as identified in the figure captions. After using these tools, the authors reviewed and edited the content as needed and take full responsibility for the content of the published article.'),
('Data availability','All input data used in this study are publicly available or reported in full in the article. The electricity prices are the publicly available Low Carbon Contracts Company IMRP actuals; the adopted workload profiles, deferral tranches and model parameters are given in Tables I, III and IV and Supplementary Section S1. The results supporting the findings are reported in the article and its supplementary information.'),
('Author confirmations','The corresponding author confirmed the author order, corresponding-author role, approval by all authors, originality/exclusive submission, no competing interests, funding completeness and funder role. No additional reviewer reports were received.'),
('Portal-specific information','Use existing author accounts/ORCIDs where available. Coauthor emails, ORCID identifiers, reviewer nominations, publishing route and any institutional agreement must come from the authors or their existing account records; they have not been invented here. Suggested reviewers are only needed if the portal requires them.')
]:
    doc.add_heading(heading,1); paragraphs(doc,text)
doc.add_heading('Abstract',1)
source=(PAPER/'main.tex').read_text(encoding='utf8')
abstract=source.split(r'\begin{abstract}',1)[1].split(r'\end{abstract}',1)[0].strip()
abstract=abstract.replace('~',' ').replace(r'\%','%').replace('--','–')
doc.add_paragraph(abstract)
doc.save(OUT/'submission_metadata.docx')
print('Created cover letter, highlights, response and submission metadata DOCX files.')
